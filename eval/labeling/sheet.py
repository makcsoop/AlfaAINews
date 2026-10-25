"""Excel-файл разметчика: запись пар и чтение заполненных меток (ТЗ «разметка 60 пар»).

Лист «Разметка» — одна строка на пару: серые колонки заполнены, жёлтые заполняет человек. Лист «Каналы» —
список для выпадающего меню, «Шкала» — памятка. Скрытый лист «_meta» хранит story_id и company_id каждой пары
и имя разметчика: импорту не нужна база, чтобы собрать CSV для метрик.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CHANNELS = Path(__file__).resolve().parents[2] / "config" / "channels_matrix.yaml"
SHEET = "Разметка"
META = "_meta"
GREY = PatternFill("solid", fgColor="F0EFEC")
YELLOW = PatternFill("solid", fgColor="FFF2B3")

# (заголовок, ключ, ширина); ключ в скобках заголовка — по нему импорт находит колонку
GIVEN = [
    ("№ пары (pair_id)", "pair_id", 11),
    ("Компания", "company", 22),
    ("Роль читателя", "role", 14),
    ("Профиль компании", "profile", 60),
    ("Дата", "date", 11),
    ("Новость", "story", 70),
    ("Источники", "sources", 40),
]
FILLED = [
    ("Значимость 0–3 (relevance)", "relevance", 12),
    ("Главный канал (main_channel)", "main_channel", 30),
    ("Знак −1/0/+1 (sign)", "sign", 10),
    ("Комментарий (comment)", "comment", 40),
]
SCALE_NOTE = [
    ("3", "важно сегодня", "нужно действовать или решать в ближайшие дни; прямо задевает деньги компании"),
    ("2", "важно", "влияет на планы, бюджет или договоры в ближайшие недели и месяцы"),
    ("1", "фон", "полезно знать для картины рынка, действий не требует"),
    ("0", "не про нас", "на компанию не влияет или влияет ничтожно"),
]


@dataclass(frozen=True)
class PairRow:
    pair_id: str
    story_id: str
    company_id: str
    company: str
    role: str
    profile: str
    date: str
    story: str
    sources: str
    item_ids: tuple[str, ...] = ()  # материалы сюжета: по ним метрики найдут пару после пересклейки


@dataclass
class LabelRow:
    pair_id: str
    story_id: str
    company_id: str
    labeler: str
    relevance: int
    main_channel: str | None
    sign: int | None
    comment: str | None
    item_ids: tuple[str, ...] = ()


@dataclass
class ReadResult:
    labeler: str
    labels: list[LabelRow] = field(default_factory=list)
    unfilled: list[str] = field(default_factory=list)  # pair_id без значимости
    errors: list[str] = field(default_factory=list)  # строка не принята
    warnings: list[str] = field(default_factory=list)  # принята, но нарушает ТЗ


def write_sheet(
    path: Path, labeler: str, batch: str, rows: Sequence[PairRow], channels: Mapping[str, str]
) -> None:
    """channels: id → название, в порядке config/channels_matrix.yaml."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET
    head = GIVEN + FILLED
    ws.append([h for h, _, _ in head])
    for col, (_, key, width) in enumerate(head, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = YELLOW if key in {k for _, k, _ in FILLED} else GREY
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    for r in rows:
        ws.append([r.pair_id, r.company, r.role, r.profile, r.date, r.story, r.sources])
    n = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=n, max_col=len(head)):
        for i, cell in enumerate(row):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = GREY if i < len(GIVEN) else YELLOW
    ws.freeze_panes = "B2"

    ch = wb.create_sheet("Каналы")
    ch.append(["Канал", "id"])
    for cid, name in channels.items():
        ch.append([name, cid])
    ch.column_dimensions["A"].width = 50
    ch.column_dimensions["B"].width = 24

    note = wb.create_sheet("Шкала")
    note.append(["Оценка", "Смысл", "Как узнать"])
    for line in SCALE_NOTE:
        note.append(list(line))
    note.column_dimensions["C"].width = 90

    base = len(GIVEN)
    rel = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    sign = DataValidation(type="list", formula1='"-1,0,1"', allow_blank=True)
    chan = DataValidation(type="list", formula1=f"='Каналы'!$A$2:$A${len(channels) + 1}", allow_blank=True)
    for dv, col in ((rel, base + 1), (chan, base + 2), (sign, base + 3)):
        letter = get_column_letter(col)
        dv.add(f"{letter}2:{letter}{n}")
        ws.add_data_validation(dv)

    meta = wb.create_sheet(META)
    meta.append(["labeler", labeler])
    meta.append(["batch", batch])
    meta.append(["pair_id", "story_id", "company_id", "item_ids"])
    for r in rows:
        meta.append([r.pair_id, r.story_id, r.company_id, " ".join(r.item_ids)])
    meta.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def load_channels(path: Path = CHANNELS) -> dict[str, str]:
    """id канала → название, в порядке config/channels_matrix.yaml."""
    return {k: v["name"] for k, v in yaml.safe_load(path.read_text(encoding="utf-8"))["channels"].items()}


def _int(v: object) -> int | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return int(str(v).strip().replace("−", "-").replace("+", ""))


def _text(v: object) -> str | None:
    s = "" if v is None else str(v).strip()
    return s or None


def read_meta(path: Path) -> tuple[str, dict[str, tuple[str, str, tuple[str, ...]]]]:
    """Разметчик и pair_id → (story_id, company_id, материалы) из скрытого листа выданного файла."""
    meta = load_workbook(path, data_only=True)[META]
    ids = {
        str(r[0]): (str(r[1]), str(r[2]), tuple(str(r[3] or "").split()) if len(r) > 3 else ())
        for r in meta.iter_rows(min_row=4, values_only=True)
        if r[0] is not None
    }
    return str(meta["B1"].value or path.stem).strip(), ids


def read_sheet(
    path: Path,
    channels: Mapping[str, str],
    origin: Path | None = None,
    labeler: str | None = None,
) -> ReadResult:
    """Метки из заполненного файла. Канал принимается названием из списка или id.

    Если разметчик пересохранил таблицу без листа «_meta» (Google Таблицы, копия листа), связки пар берутся
    из выданного файла origin. Колонки ищутся по ключу в скобках заголовка или по самому ключу.
    """
    wb = load_workbook(path, data_only=True)
    if META in wb.sheetnames:
        name, ids = read_meta(path)
    elif origin is not None:
        name, ids = read_meta(origin)
    else:
        raise ValueError(f"{path.name}: нет листа {META} — укажите выданный файл (--origin)")
    labeler = labeler or name
    by_name = {name.strip().lower(): cid for cid, name in channels.items()}
    by_name.update({cid: cid for cid in channels})
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    header = [str(c.value or "").strip() for c in ws[1]]
    col = {key: next(i for i, h in enumerate(header) if f"({key})" in h or h == key) for key in _keys()}
    out = ReadResult(labeler=labeler)
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        pid = _text(row[col["pair_id"]])
        if pid is None:
            continue
        where = f"{path.name}, строка {n} ({pid})"
        if pid not in ids:
            out.errors.append(f"{where}: такой пары нет в файле")
            continue
        try:
            rel = _int(row[col["relevance"]])
            sign = _int(row[col["sign"]])
        except ValueError:
            out.errors.append(f"{where}: значимость и знак — целые числа")
            continue
        if rel is None:
            out.unfilled.append(pid)
            continue
        if rel not in (0, 1, 2, 3):
            out.errors.append(f"{where}: значимость {rel} вне шкалы 0–3")
            continue
        if sign is not None and sign not in (-1, 0, 1):
            out.errors.append(f"{where}: знак {sign} — только −1, 0 или +1")
            continue
        raw_channel = _text(row[col["main_channel"]])
        channel = by_name.get(raw_channel.lower()) if raw_channel else None
        if raw_channel and channel is None:
            out.errors.append(f"{where}: канала «{raw_channel}» нет в списке")
            continue
        comment = _text(row[col["comment"]])
        if rel >= 1 and (channel is None or sign is None):
            out.warnings.append(f"{where}: при значимости ≥ 1 нужны канал и знак")
        if rel >= 2 and comment is None:
            out.warnings.append(f"{where}: при значимости 2–3 нужен комментарий")
        sid, cid, items = ids[pid]
        out.labels.append(LabelRow(pid, sid, cid, labeler, rel, channel, sign, comment, items))
    return out


def _keys() -> Iterable[str]:
    return ("pair_id", "relevance", "main_channel", "sign", "comment")
