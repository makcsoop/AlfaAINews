"""Разметка 60 пар: отбор пулом, раздача разметчикам, Excel туда и обратно, CSV для метрик."""

import random

from openpyxl import Workbook, load_workbook

from backend.eval.agreement import read_labels
from backend.eval.metrics import gold_from_csv
from backend.schemas import CompanyProfile
from eval.labeling.import_labels import write_csv
from eval.labeling.make_pairs import assign, pool, profile_text, sources_text
from eval.labeling.sheet import SHEET, PairRow, read_sheet, write_sheet

CHANNELS = {"key_rate": "Ключевая ставка и стоимость денег", "fx_rub": "Курс рубля"}


def test_pool_takes_both_tops_and_random_without_repeats():
    system = {f"s{i}": 1 - i / 100 for i in range(30)}  # s0 лучший
    baseline = {f"s{i}": i / 100 for i in range(30)}  # s29 лучший
    got = pool(system, baseline, random.Random(1))
    ids = [s for s, _ in got]
    assert ids[:4] == ["s0", "s1", "s2", "s3"]
    assert ids[4:7] == ["s29", "s28", "s27"]
    assert len(set(ids)) == 10 and [w for _, w in got].count("random") == 3


def test_pool_baseline_skips_system_picks():
    same = {f"s{i}": 1 - i / 100 for i in range(12)}
    ids = [s for s, _ in pool(same, same, random.Random(1))]
    assert ids[4:7] == ["s4", "s5", "s6"] and len(set(ids)) == 10


def test_assign_common_pairs_for_everyone_rest_split():
    groups = {f"c{c}": [f"c{c}-{i}" for i in range(10)] for c in range(6)}
    per, common = assign(groups, ["a", "b", "c"], random.Random(1))
    assert len(common) == 10
    assert {g.split("-")[0] for g in common} == set(groups)  # общие — от каждого клиента
    for name in per:
        assert common <= set(per[name])
    own = [p for name in per for p in per[name] if p not in common]
    assert sorted(own) == sorted(p for g in groups.values() for p in g if p not in common)  # каждая ровно раз
    assert max(len(v) for v in per.values()) - min(len(v) for v in per.values()) <= 1


def test_single_labeler_gets_everything():
    groups = {"c": ["p1", "p2"]}
    per, common = assign(groups, ["solo"], random.Random(1))
    assert per == {"solo": ["p1", "p2"]} and common == set()


def test_profile_text_human_readable():
    p = CompanyProfile(
        company_id="U_1",
        okved_code="46.71",
        okved_name="Торговля оптовая топливом",
        revenue=60.4e9,
        revenue_rank_in_okved=3,
        revenue_peers_in_okved=120,
        region="г. Москва",
        fts_flag=True,
        fts_country="Китай",
        products={"loans": True, "leasing": True, "deposits": False},
    )
    text = profile_text(
        p, {"sector_shares_suppliers": {"46": 0.7, "outside": 0.3}, "top_supplier_share": 0.35}
    )
    assert "Выручка 60,4 млрд ₽ (3-я из 120 в отрасли)" in text
    assert "Продукты в банке: кредиты, лизинг" in text and "депозиты" not in text
    assert "ВЭД: да (Китай)" in text
    assert "Поставщики: оптовая торговля: 70%, вне базы: 30%; на крупнейшего — 35%" in text


def test_sources_one_link_per_outlet():
    items = [
        {"source_id": "cbr", "url": "https://cbr.ru/1"},
        {"source_id": "cbr", "url": "https://cbr.ru/2"},
        {"source_id": "tass", "url": None},
        {"source_id": "rbc", "url": "https://rbc.ru/x"},
    ]
    assert (
        sources_text(items, {"cbr": "Банк России"}) == "Банк России: https://cbr.ru/1\nrbc: https://rbc.ru/x"
    )


def _rows(n=4):
    return [
        PairRow(
            f"oct-60-0{i}",
            f"s{i}",
            "U_1",
            "Оптовик (U_1)",
            "финдиректор",
            "профиль",
            "23.10.2026",
            "новость",
            "",
        )
        for i in range(1, n + 1)
    ]


def _fill(path, values):
    wb = load_workbook(path)
    ws = wb[SHEET]
    for r, vals in enumerate(values, start=2):
        for c, v in enumerate(vals, start=8):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


def test_sheet_roundtrip_validates(tmp_path):
    path = tmp_path / "razmetka_r1.xlsx"
    write_sheet(path, "r1", "oct-60", _rows(), CHANNELS)
    wb = load_workbook(path)
    assert wb["_meta"].sheet_state == "hidden" and wb[SHEET].max_row == 5
    _fill(
        path,
        [
            (3, "Ключевая ставка и стоимость денег", -1, "кредит по плавающей"),  # всё верно
            (2, "fx_rub", "+1", None),  # id канала, «+1», нет комментария → замечание
            (5, None, None, None),  # вне шкалы → ошибка
            (None, None, None, None),  # не размечено
        ],
    )
    r = read_sheet(path, CHANNELS)
    assert r.labeler == "r1"
    assert [(x.pair_id, x.relevance, x.main_channel, x.sign) for x in r.labels] == [
        ("oct-60-01", 3, "key_rate", -1),
        ("oct-60-02", 2, "fx_rub", 1),
    ]
    assert r.labels[0].story_id == "s1" and r.labels[0].company_id == "U_1"
    assert len(r.errors) == 1 and "вне шкалы" in r.errors[0]
    assert r.unfilled == ["oct-60-04"]
    assert any("комментарий" in w for w in r.warnings)


def test_unknown_channel_is_error(tmp_path):
    path = tmp_path / "f.xlsx"
    write_sheet(path, "r2", "oct-60", _rows(1), CHANNELS)
    _fill(path, [(1, "Погода", 0, None)])
    r = read_sheet(path, CHANNELS)
    assert not r.labels and "Погода" in r.errors[0]


def test_resaved_sheet_without_meta_uses_origin(tmp_path):
    # разметчик пересохранил таблицу: один лист, заголовок «pair_id», числа дробями, каналы по id
    origin = tmp_path / "razmetka.xlsx"
    write_sheet(origin, "razmetka", "oct-60", _rows(2), CHANNELS)
    wb = Workbook()
    ws = wb.active
    ws.title = "Untitled"
    ws.append(
        [
            "pair_id",
            "Компания",
            "Значимость 0–3 (relevance)",
            "Главный канал (main_channel)",
            "Знак −1/0/+1 (sign)",
            "Комментарий (comment)",
        ]
    )
    ws.append(["oct-60-01", "Оптовик", 2.0, "key_rate", -1.0, "кредит"])
    ws.append(["oct-60-02", "Оптовик", 0.0, None, 0.0, None])
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    r = read_sheet(path, CHANNELS, origin)
    assert r.labeler == "razmetka" and not r.errors and not r.warnings
    assert [(x.pair_id, x.story_id, x.relevance, x.main_channel, x.sign) for x in r.labels] == [
        ("oct-60-01", "s1", 2, "key_rate", -1),
        ("oct-60-02", "s2", 0, None, 0),
    ]


def test_csv_feeds_metrics_and_agreement(tmp_path):
    path = tmp_path / "x.xlsx"
    write_sheet(path, "r1", "oct-60", _rows(2), CHANNELS)
    _fill(path, [(2, "fx_rub", 1, "да"), (1, "fx_rub", 0, None)])
    out = tmp_path / "labels.csv"
    write_csv(out, read_sheet(path, CHANNELS).labels)
    assert gold_from_csv(out) == {"oct-60-01": ("s1", "U_1", 1.0), "oct-60-02": ("s2", "U_1", 0.0)}
    assert [(x.labeler, x.relevance) for x in read_labels(out)] == [("r1", 2), ("r1", 1)]
